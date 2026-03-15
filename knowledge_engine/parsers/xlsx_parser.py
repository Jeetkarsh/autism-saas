"""XLSX parser using openpyxl."""
import io
from parsers.base import BaseParser

class XlsxParser(BaseParser):
    """Parser for .xlsx files using openpyxl."""

    def parse(self, content: bytes, filename: str) -> str:
        import openpyxl  # Wait to import until execution

        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        text_parts = []
        
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            text_parts.append(f"--- Sheet: {sheet_name} ---")
            
            # Extract column headers from the first row
            headers = []
            for col in range(1, sheet.max_column + 1):
                cell_value = sheet.cell(row=1, column=col).value
                headers.append(str(cell_value) if cell_value is not None else f"Column {col}")
                
            # Process data rows
            for row_idx in range(2, sheet.max_row + 1):
                row_data = []
                has_data = False
                for col_idx in range(1, sheet.max_column + 1):
                    cell_value = sheet.cell(row=row_idx, column=col_idx).value
                    if cell_value is not None:
                        has_data = True
                        val_str = str(cell_value).strip()
                        if val_str:
                            row_data.append(f"{headers[col_idx-1]}: {val_str}")
                            
                if has_data and row_data:
                    text_parts.append(" | ".join(row_data))
                    
            text_parts.append("\n")
            
        return "\n".join(text_parts)

    def supported_extensions(self) -> list[str]:
        return [".xlsx"]
