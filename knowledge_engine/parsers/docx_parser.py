"""DOCX parser using python-docx."""
import io
from parsers.base import BaseParser

class DocxParser(BaseParser):
    """Parser for .docx files using python-docx."""

    def parse(self, content: bytes, filename: str) -> str:
        import docx  # Wait to import until execution

        doc = docx.Document(io.BytesIO(content))
        text_parts = []
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                text_parts.append(text)
        
        # Also extract text from tables if there are any
        for table in doc.tables:
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_data:
                    text_parts.append(" | ".join(row_data))
                    
        return "\n\n".join(text_parts)

    def supported_extensions(self) -> list[str]:
        return [".docx"]
